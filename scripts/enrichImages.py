"""Reaplica fotos verificadas sem alterar nenhum outro dado (Python + openpyxl)."""
import json
from pathlib import Path
from openpyxl import load_workbook

project = Path(__file__).resolve().parent.parent
manifest = json.loads((project / 'data-source/image-sources.json').read_text(encoding='utf-8'))
original = project / 'data-source' / manifest['originalWorkbook']
destination = project / 'data-source' / manifest['enrichedWorkbook']
book = load_workbook(original)
sheet = book['Importar_Olist']
columns = {c.value: c.column for c in sheet[1]}
image_columns = [n for name, n in columns.items() if str(name).startswith('URL imagem')]
rows = list(sheet.iter_rows(min_row=2))
def value(row, key):
    return row[columns[key]-1].value
for entry in manifest['entries']:
    parent = next((r for r in rows if value(r, 'Tipo do produto') == 'V' and str(value(r, 'Cód do fornecedor')) == entry['reference']), None)
    if parent is None:
        raise ValueError(f"Referência não encontrada: {entry['reference']}")
    def matches(row):
        parts = dict(p.split(':', 1) for p in str(value(row, 'Variações') or '').split('||') if ':' in p)
        return value(row, 'Código do pai') == value(parent, 'Código (SKU)') and parts.get('Cor', '') == entry['color']
    targets = [r for r in rows if matches(r)] if entry['color'] else [parent]
    if not targets:
        raise ValueError(f"Cor não encontrada: {entry['reference']} / {entry['color']}")
    for row in targets:
        for url in entry['imageUrls']:
            if url in [row[n-1].value for n in image_columns]:
                continue
            empty = next((n for n in image_columns if not row[n-1].value), None)
            if empty is None:
                break  # Limite de colunas Olist; preservar links originais.
            row[empty-1].value = url
book.save(destination)
before, after = load_workbook(original), load_workbook(destination)
changes = []
assert before.sheetnames == after.sheetnames
for name in before.sheetnames:
    a, b = before[name], after[name]
    assert (a.max_row, a.max_column) == (b.max_row, b.max_column)
    for original_row, enriched_row in zip(a.iter_rows(), b.iter_rows()):
        for old, new in zip(original_row, enriched_row):
            if old.value == new.value:
                continue
            assert name == 'Importar_Olist' and old.column in image_columns and not old.value
            changes.append({'sheet': name, 'cell': old.coordinate, 'url': new.value})
audit = {'originalWorkbook': manifest['originalWorkbook'], 'enrichedWorkbook': manifest['enrichedWorkbook'], 'nonImageCellsChanged': 0, 'existingImageLinksChanged': 0, 'imageCellsAdded': len(changes), 'changes': changes}
(project / 'docs/enriquecimento-imagens.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
print(f"{len(changes)} células vazias de imagem preenchidas. Demais células preservadas.")
